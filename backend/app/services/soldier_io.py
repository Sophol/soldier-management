"""Excel import, export and comparison of professional soldiers (New.txt §2 ការងារបញ្ចូលក្របខ័ណ្ឌយោធិនអាជីព,
§3 ការប្រៀបធៀប). One field catalog drives the template, the import, the export and the comparison, so a file
exported from the system can be imported or compared again as it is."""

import re
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import Lookup, Position, Rank, Soldier, Specialty, Unit
from app.services.soldier import unit_chain


@dataclass(frozen=True)
class Field:
    key: str  # Soldier attribute
    label: str  # Excel column header
    kind: str  # text / date / int / gender / track / rank / position / unit / specialty / lookup:<category>
    required: bool = False


FIELDS = [
    Field("military_id", "អត្តលេខ", "text", True),
    Field("name_kh", "គោត្តនាមនិងនាម", "text", True),
    Field("name_latin", "ឈ្មោះជាអក្សរឡាតាំង", "text"),
    Field("gender", "ភេទ", "gender", True),
    Field("date_of_birth", "ថ្ងៃខែឆ្នាំកំណើត", "date"),
    Field("national_id", "លេខអត្តសញ្ញាណប័ណ្ណសញ្ជាតិខ្មែរ", "text"),
    Field("national_id_expiry", "ថ្ងៃផុតកំណត់អត្តសញ្ញាណប័ណ្ណ", "date"),
    Field("nationality", "សញ្ជាតិ", "text"),
    Field("ethnicity", "ជនជាតិ", "text"),
    Field("phone", "លេខទូរសព្ទ", "text"),
    Field("current_rank_id", "ឋានន្តរស័ក្តិ", "rank"),
    Field("rank_date", "ថ្ងៃខែឆ្នាំឡើងស័ក្តិ", "date"),
    Field("current_position_id", "មុខតំណែង", "position"),
    Field("position_date", "ថ្ងៃខែឆ្នាំមុខតំណែង", "date"),
    Field("current_unit_id", "ស្ថាប័នកងឯកភាព-អង្គភាព", "unit"),
    Field("unit_transfer_date", "ថ្ងៃខែឆ្នាំផ្លាស់ប្ដូរកងឯកភាព", "date"),
    Field("from_unit_text", "មកពីស្ថាប័នកងឯកភាព", "text"),
    Field("enlistment_date", "ថ្ងៃខែឆ្នាំបម្រើចូលទ័ព", "date"),
    Field("establishment_no", "បញ្ចូលក្របខ័ណ្ឌលេខ", "text"),
    Field("service_track", "ប្រភេទសេវា", "track"),
    Field("status_id", "ស្ថានភាព", "lookup:soldier_status"),
    Field("education_level_id", "កម្រិតវប្បធម៌", "lookup:education_level"),
    Field("specialty_id", "ជំនាញឯកទេស", "specialty"),
    Field("blood_type_id", "ប្រភេទឈាម", "lookup:blood_type"),
    Field("marital_status_id", "ស្ថានភាពគ្រួសារ", "lookup:marital_status"),
    Field("spouse_name", "ឈ្មោះប្ដីឬប្រពន្ធ", "text"),
    Field("spouse_job_id", "មុខរបរប្ដីឬប្រពន្ធ", "lookup:occupation"),
    Field("children_count", "ចំនួនកូន", "int"),
    Field("father_name", "ឈ្មោះឪពុក", "text"),
    Field("mother_name", "ឈ្មោះម្ដាយ", "text"),
    Field("birth_place_text", "ទីកន្លែងកំណើត", "text"),
    Field("current_address_text", "ទីលំនៅបច្ចុប្បន្ន", "text"),
    Field("voting_place", "ទីបោះឆ្នោត", "text"),
    Field("dependents_count", "ចំនួនអ្នកក្នុងបន្ទុក", "int"),
    Field("notes", "ប្រវត្តិផ្សេងៗ", "text"),
]
BY_KEY = {f.key: f for f in FIELDS}
SERVICE_FIELDS = ("current_rank_id", "current_position_id", "current_unit_id")
GENDERS = {"ប្រុស": "M", "ស្រី": "F", "m": "M", "f": "F", "male": "M", "female": "F"}
TRACKS = {"សេវារដ្ឋបាល": "admin", "រដ្ឋបាល": "admin", "admin": "admin",
          "ប្រតិបត្តិការសឹក": "combat", "សឹក": "combat", "combat": "combat"}
KH_DIGITS = str.maketrans("០១២៣៤៥៦៧៨៩", "0123456789")
MAX_ROWS = 5000


def _norm(text) -> str:
    """Trimmed, single-spaced text (Khmer digits kept: they are part of names such as លេខ១)."""
    return re.sub(r"\s+", " ", str(text or "")).strip()


def _digits(text) -> str:
    """For dates and numbers only: Khmer digits read as 0-9."""
    return _norm(text).translate(KH_DIGITS)


def _header_key(text) -> str:
    return _norm(text).replace("*", "").strip().lower()


class Resolver:
    """Turns the text people type in Excel into ids (and back), with the reason when it cannot."""

    def __init__(self, db: Session):
        self.db = db
        self.ranks = list(db.scalars(select(Rank)))
        self.positions = list(db.scalars(select(Position)))
        self.specialties = list(db.scalars(select(Specialty)))
        self.units = list(db.scalars(select(Unit)))
        self.lookups: dict[str, list[Lookup]] = {}
        for lk in db.scalars(select(Lookup)):
            self.lookups.setdefault(lk.category, []).append(lk)
        self.unit_path = {u.id: " › ".join(x.name_kh for x in unit_chain(u)) for u in self.units}

    @staticmethod
    def _one(matches: list, value: str, what: str):
        if len(matches) == 1:
            return matches[0].id, None
        if not matches:
            return None, f"រកមិនឃើញ{what} '{value}'"
        return None, f"{what} '{value}' មានច្រើនជាងមួយ — សូមសរសេរឲ្យច្បាស់ (ឧ. ផ្លូវពេញរបស់អង្គភាព ឬកូដ)"

    def parse(self, f: Field, raw):
        """(value, error) for one cell; (None, None) for an empty cell."""
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            return None, None
        text = _norm(raw)
        if f.kind == "text":
            return (str(int(raw)) if isinstance(raw, float) and raw.is_integer() else text), None
        if f.kind == "int":
            try:
                return int(float(_digits(raw))), None
            except ValueError:
                return None, f"{f.label}: '{raw}' មិនមែនជាលេខ"
        if f.kind == "date":
            if isinstance(raw, datetime):
                return raw.date(), None
            if isinstance(raw, date):
                return raw, None
            for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
                try:
                    return datetime.strptime(_digits(raw), fmt).date(), None
                except ValueError:
                    continue
            return None, f"{f.label}: '{raw}' មិនមែនជាថ្ងៃខែឆ្នាំ (ថ្ងៃ/ខែ/ឆ្នាំ)"
        if f.kind == "gender":
            g = GENDERS.get(text.lower())
            return (g, None) if g else (None, f"{f.label}: '{raw}' (ប្រុស ឬ ស្រី)")
        if f.kind == "track":
            t = TRACKS.get(text.lower())
            return (t, None) if t else (None, f"{f.label}: '{raw}' (សេវារដ្ឋបាល ឬ ប្រតិបត្តិការសឹក)")
        names = lambda *values: {_norm(v) for v in values if v}  # noqa: E731
        if f.kind == "rank":
            return self._one([r for r in self.ranks if text in names(r.name_kh, r.short_kh, r.name_en)], text, "ឋានន្តរស័ក្តិ")
        if f.kind == "position":
            return self._one([p for p in self.positions if text in names(p.name_kh, p.name_en)], text, "មុខតំណែង")
        if f.kind == "specialty":
            return self._one([s for s in self.specialties if text in names(s.name_kh, s.name_en)], text, "ជំនាញឯកទេស")
        if f.kind == "unit":
            path = re.sub(r"\s*[›>/]\s*", " › ", text)
            by_path = [u for u in self.units if _norm(self.unit_path[u.id]) == path or u.code == text]
            return self._one(by_path or [u for u in self.units if _norm(u.name_kh) == text], text, "អង្គភាព")
        if f.kind.startswith("lookup:"):
            items = self.lookups.get(f.kind[7:], [])
            return self._one([lk for lk in items if text in names(lk.name_kh, lk.name_en, lk.code)], text, f.label)
        return text, None

    def display(self, f: Field, value) -> str:
        """The value as it is written in the Excel file (names, dd/mm/yyyy)."""
        if value is None or value == "":
            return ""
        if f.kind == "date":
            return value.strftime("%d/%m/%Y")
        if f.kind == "gender":
            return "ប្រុស" if value == "M" else "ស្រី"
        if f.kind == "track":
            return "ប្រតិបត្តិការសឹក" if value == "combat" else "សេវារដ្ឋបាល"
        if f.kind == "unit":
            return self.unit_path.get(value, "")
        pools = {"rank": self.ranks, "position": self.positions, "specialty": self.specialties}
        if f.kind in pools:
            return next((x.name_kh for x in pools[f.kind] if x.id == value), "")
        if f.kind.startswith("lookup:"):
            return next((x.name_kh for x in self.lookups.get(f.kind[7:], []) if x.id == value), "")
        return str(value)


# ---------------- reading files ----------------

def read_sheet(content: bytes) -> tuple[list[Field], list[tuple[int, dict]], list[str]]:
    """(fields found in the header, [(excel row number, {key: raw cell})], unknown headers)."""
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 - any unreadable file gets the same answer
        raise HTTPException(400, "មិនអាចអានឯកសារ Excel (.xlsx) នេះបានទេ") from e
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    labels = {_header_key(f.label): f for f in FIELDS} | {f.key: f for f in FIELDS}
    header, header_row = None, 0
    for header_row, row in enumerate(rows, start=1):  # the header may sit below a title row
        if any(_header_key(c) in labels for c in row if c):
            header = row
            break
        if header_row >= 10:
            break
    if not header:
        raise HTTPException(400, "រកមិនឃើញជួរចំណងជើង (ត្រូវមានយ៉ាងហោចណាស់ 'អត្តលេខ')")
    columns = {i: labels.get(_header_key(c)) for i, c in enumerate(header)}
    fields = [f for f in columns.values() if f]
    if BY_KEY["military_id"] not in fields:
        raise HTTPException(400, "ឯកសារត្រូវមានជួរ 'អត្តលេខ'")
    unknown = [str(c) for i, c in enumerate(header) if c and not columns[i]]
    data = []
    for n, row in enumerate(rows, start=header_row + 1):
        if not any(c not in (None, "") for c in row):
            continue
        data.append((n, {columns[i].key: v for i, v in enumerate(row) if i in columns and columns[i]}))
        if len(data) > MAX_ROWS:
            raise HTTPException(400, f"ឯកសារមានជួរច្រើនពេក (អតិបរមា {MAX_ROWS})")
    return fields, data, unknown


# ---------------- template & export ----------------

def _book(title: str, fields: list[Field]) -> tuple[Workbook, object]:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append([f"{f.label} *" if f.required else f.label for f in fields])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8F0EC")
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 22
    ws.freeze_panes = "A2"
    return wb, ws


def _save(wb: Workbook) -> BytesIO:
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def template(db: Session) -> BytesIO:
    """Empty import sheet plus a sheet of the accepted values for every list column."""
    r = Resolver(db)
    wb, ws = _book("យោធិន", FIELDS)
    help_ws = wb.create_sheet("តម្លៃដែលអាចប្រើ")
    lists = [
        ("ឋានន្តរស័ក្តិ", [x.name_kh for x in sorted(r.ranks, key=lambda x: -x.sort_order) if x.active]),
        ("មុខតំណែង", [x.name_kh for x in r.positions if x.active]),
        ("ស្ថាប័នកងឯកភាព-អង្គភាព (ផ្លូវពេញ ឬកូដ)", sorted(r.unit_path.values())),
        ("ស្ថានភាព", [x.name_kh for x in r.lookups.get("soldier_status", [])]),
        ("កម្រិតវប្បធម៌", [x.name_kh for x in r.lookups.get("education_level", [])]),
        ("ជំនាញឯកទេស", [x.name_kh for x in r.specialties if x.active]),
        ("ប្រភេទឈាម", [x.name_kh for x in r.lookups.get("blood_type", [])]),
        ("ស្ថានភាពគ្រួសារ", [x.name_kh for x in r.lookups.get("marital_status", [])]),
        ("មុខរបរ", [x.name_kh for x in r.lookups.get("occupation", [])]),
        ("ភេទ / ប្រភេទសេវា", ["ប្រុស", "ស្រី", "សេវារដ្ឋបាល", "ប្រតិបត្តិការសឹក"]),
    ]
    help_ws.append([title for title, _ in lists])
    for c in help_ws[1]:
        c.font = Font(bold=True)
    for i in range(max(len(v) for _, v in lists)):
        help_ws.append([v[i] if i < len(v) else None for _, v in lists])
    for col in help_ws.columns:
        help_ws.column_dimensions[col[0].column_letter].width = 30
    return _save(wb)


EXTRA_EXPORT = [("age", "អាយុ"), ("service_length", "រយៈពេលបម្រើការងារ"), ("time_in_rank", "រយៈពេលឡើងស័ក្តិ"),
                ("time_in_position", "រយៈពេលមុខតំណែង"), ("institution_name", "ចំណុះស្ថាប័នកងឯកភាព")]


def export(db: Session, soldiers, summary) -> BytesIO:
    """The import columns (so the file can be imported/compared again) followed by computed columns."""
    r = Resolver(db)
    wb, ws = _book("យោធិន", FIELDS + [Field(k, label, "text") for k, label in EXTRA_EXPORT])
    for s in soldiers:
        d = summary(s)
        ws.append([r.display(f, getattr(s, f.key)) for f in FIELDS] + [d.get(k) for k, _ in EXTRA_EXPORT])
    return _save(wb)


# ---------------- import ----------------

def parse_rows(db: Session, data: list[tuple[int, dict]], fields: list[Field]) -> list[dict]:
    """Every row with its parsed values and errors; nothing is written."""
    r = Resolver(db)
    seen_mid, seen_nid, out = {}, {}, []
    for n, raw in data:
        values, errors, invalid = {}, [], set()
        for f in fields:
            v, err = r.parse(f, raw.get(f.key))
            if err:
                errors.append(err)
                invalid.add(f.key)
            elif v is not None:
                values[f.key] = v
        mid = values.get("military_id")
        if not mid:
            errors.append("ខ្វះអត្តលេខ")
        elif mid in seen_mid:
            errors.append(f"អត្តលេខស្ទួនក្នុងឯកសារ (ជួរ {seen_mid[mid]})")
        else:
            seen_mid[mid] = n
        nid = values.get("national_id")
        if nid and nid in seen_nid:
            errors.append(f"លេខអត្តសញ្ញាណប័ណ្ណស្ទួនក្នុងឯកសារ (ជួរ {seen_nid[nid]})")
        elif nid:
            seen_nid[nid] = n
        out.append({"row": n, "values": values, "errors": errors, "invalid": invalid})
    return out


def plan_import(db: Session, rows: list[dict], update_existing: bool, scope: list[int] | None) -> list[dict]:
    """Decide create / update / error for each parsed row (still nothing written)."""
    mids = [r["values"]["military_id"] for r in rows if r["values"].get("military_id")]
    existing = {s.military_id: s for s in db.scalars(select(Soldier).where(Soldier.military_id.in_(mids)))} if mids else {}
    nids = [r["values"]["national_id"] for r in rows if r["values"].get("national_id")]
    taken_nid = dict(db.execute(select(Soldier.national_id, Soldier.military_id).where(
        Soldier.national_id.in_(nids))).all()) if nids else {}
    for r in rows:
        v, errs = r["values"], r["errors"]
        s = existing.get(v.get("military_id"))
        r["name_kh"] = v.get("name_kh") or (s.name_kh if s else None)
        if s is not None:
            if s.deleted_at:
                errs.append("អត្តលេខនេះជារបស់យោធិនដែលត្រូវបានលុប")
            elif not update_existing:
                errs.append("អត្តលេខនេះមានក្នុងប្រព័ន្ធរួចហើយ")
            r["action"], r["soldier_id"] = "update", s.id
        else:
            for f in FIELDS:
                if f.required and not v.get(f.key) and f.key not in ("military_id", *r["invalid"]):
                    errs.append(f"ខ្វះ {f.label}")
            r["action"] = "create"
        nid = v.get("national_id")
        if nid and nid in taken_nid and taken_nid[nid] != v.get("military_id"):
            errs.append(f"លេខអត្តសញ្ញាណប័ណ្ណនេះជារបស់យោធិន {taken_nid[nid]}")
        unit = v.get("current_unit_id") or (s.current_unit_id if s else None)
        if scope is not None and unit not in scope:
            errs.append("អង្គភាពនៅក្រៅសិទ្ធិរបស់អ្នក")
        if errs:
            r["action"] = "error"
    return rows


def import_summary(rows: list[dict]) -> dict:
    count = {a: sum(1 for r in rows if r["action"] == a) for a in ("create", "update", "error")}
    return {**count, "total": len(rows),
            "rows": [{k: r.get(k) for k in ("row", "action", "errors", "name_kh", "soldier_id")}
                     | {"military_id": r["values"].get("military_id")} for r in rows]}


# ---------------- comparison ----------------

def compare(db: Session, data: list[tuple[int, dict]], fields: list[Field], soldiers_in_scope) -> dict:
    """Compare the file with the soldier list on the chosen fields, matching rows by military ID.
    soldiers_in_scope: the soldiers the file is expected to cover (for 'missing from the file')."""
    r = Resolver(db)
    by_mid = {s.military_id: s for s in soldiers_in_scope}
    file_mids = set()
    out = []
    for n, raw in data:
        mid = _norm(raw.get("military_id"))
        if not mid:
            continue
        file_mids.add(mid)
        s = by_mid.get(mid) or db.scalar(select(Soldier).where(func.lower(Soldier.military_id) == mid.lower(),
                                                               Soldier.deleted_at.is_(None)))
        if s is None:
            out.append({"row": n, "military_id": mid, "name_kh": _norm(raw.get("name_kh")), "status": "missing_in_db",
                        "diffs": []})
            continue
        diffs = []
        for f in fields:
            if f.key == "military_id":
                continue
            in_db = r.display(f, getattr(s, f.key))
            parsed, err = r.parse(f, raw.get(f.key))
            in_file = r.display(f, parsed) if not err else _norm(raw.get(f.key))
            if _norm(in_db) != _norm(in_file):
                diffs.append({"field": f.key, "label": f.label, "db": in_db, "file": in_file})
        out.append({"row": n, "military_id": mid, "name_kh": s.name_kh, "soldier_id": s.id,
                    "status": "different" if diffs else "same", "diffs": diffs})
    for mid, s in by_mid.items():
        if mid not in file_mids:
            out.append({"row": None, "military_id": mid, "name_kh": s.name_kh, "soldier_id": s.id,
                        "status": "missing_in_file", "diffs": []})
    count = {k: sum(1 for x in out if x["status"] == k) for k in ("same", "different", "missing_in_db", "missing_in_file")}
    return {**count, "fields": [{"key": f.key, "label": f.label} for f in fields], "rows": out}


COMPARE_STATUS = {"same": "ដូចគ្នា", "different": "ខុសគ្នា", "missing_in_db": "មិនមានក្នុងប្រព័ន្ធ",
                  "missing_in_file": "មិនមានក្នុងឯកសារ"}


def compare_workbook(result: dict) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "លទ្ធផលប្រៀបធៀប"
    ws.append(["ជួរ", "អត្តលេខ", "គោត្តនាមនិងនាម", "លទ្ធផល", "ព័ត៌មាន", "ក្នុងប្រព័ន្ធ", "ក្នុងឯកសារ"])
    for c in ws[1]:
        c.font = Font(bold=True)
    red = PatternFill("solid", fgColor="FBE3E1")
    for x in result["rows"]:
        if x["status"] == "same":
            continue
        base = [x["row"], x["military_id"], x["name_kh"], COMPARE_STATUS[x["status"]]]
        for d in x["diffs"] or [None]:
            ws.append(base + ([d["label"], d["db"], d["file"]] if d else ["", "", ""]))
            if d:
                for c in ws[ws.max_row][4:7]:
                    c.fill = red
    for col, width in zip("ABCDEFG", (8, 14, 24, 18, 26, 30, 30), strict=True):
        ws.column_dimensions[col].width = width
    return _save(wb)

"""Auth, users, roles & permissions, system settings, audit log."""

from datetime import date, datetime, timedelta
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.database import get_db
from app.core.deps import get_current_user, require
from app.core.permissions import ACTIONS, MODULES, module_actions
from app.core.security import (
    create_access_token,
    create_refresh_token,
    decode_token,
    hash_password,
    verify_password,
)
from app.models import AuditLog, Permission, Role, SavedFilter, Setting, User
from app.schemas import (
    ChangePasswordIn,
    ExportTableIn,
    LoginIn,
    RefreshIn,
    RoleIn,
    RolePermissionsIn,
    SavedFilterIn,
    SettingsIn,
    UserIn,
)
from app.services.common import audit, delete_storage_file, paginate, save_upload, storage_file, to_dict

router = APIRouter()


def user_out(u: User) -> dict:
    return {
        "id": u.id, "username": u.username, "full_name": u.full_name, "phone": u.phone,
        # the Khmer description is the label people see; the English name stays the role's key
        "role_id": u.role_id, "role_name": u.role.description or u.role.name, "is_super": u.role.is_super,
        "unit_scope_id": u.unit_scope_id, "unit_scope_name": u.unit_scope.name_kh if u.unit_scope else None,
        "active": u.active, "last_login_at": u.last_login_at.isoformat() if u.last_login_at else None,
    }


def me_out(u: User) -> dict:
    perms = sorted(f"{m}.{a}" for m in MODULES for a in module_actions(m)) if u.role.is_super else sorted(u.permission_keys())
    return {**user_out(u), "permissions": perms}


# ---------------- auth ----------------

@router.post("/auth/login", tags=["auth"])
def login(payload: LoginIn, request: Request, db: Session = Depends(get_db)):
    user = db.scalar(select(User).where(User.username == payload.username))
    now = datetime.now()
    if user and user.locked_until and user.locked_until > now:
        raise HTTPException(423, f"គណនីត្រូវបានចាក់សោបណ្ដោះអាសន្ន រហូតដល់ {user.locked_until:%H:%M}")
    if not user or not user.active or not verify_password(payload.password, user.password_hash):
        if user:
            user.failed_logins += 1
            if user.failed_logins >= settings.MAX_FAILED_LOGINS:
                user.locked_until = now + timedelta(minutes=settings.LOCKOUT_MINUTES)
                user.failed_logins = 0
            db.commit()
        raise HTTPException(401, "ឈ្មោះអ្នកប្រើ ឬពាក្យសម្ងាត់មិនត្រឹមត្រូវ")
    user.failed_logins = 0
    user.locked_until = None
    user.last_login_at = now
    audit(db, user, "login", table="users", record_id=user.id, request=request)
    db.commit()
    return {
        "access_token": create_access_token(user.id),
        "refresh_token": create_refresh_token(user.id),
        "token_type": "bearer",
        "user": me_out(user),
    }


@router.post("/auth/refresh", tags=["auth"])
def refresh(payload: RefreshIn, db: Session = Depends(get_db)):
    user_id = decode_token(payload.refresh_token, "refresh")
    user = db.get(User, user_id) if user_id else None
    if not user or not user.active:
        raise HTTPException(401, "Refresh token មិនត្រឹមត្រូវ")
    return {"access_token": create_access_token(user.id), "token_type": "bearer"}


@router.get("/auth/me", tags=["auth"])
def me(user: User = Depends(get_current_user)):
    return me_out(user)


@router.post("/auth/change-password", tags=["auth"])
def change_password(payload: ChangePasswordIn, request: Request, db: Session = Depends(get_db),
                    user: User = Depends(get_current_user)):
    if not verify_password(payload.old_password, user.password_hash):
        raise HTTPException(400, "ពាក្យសម្ងាត់ចាស់មិនត្រឹមត្រូវ")
    user.password_hash = hash_password(payload.new_password)
    audit(db, user, "change_password", table="users", record_id=user.id, request=request)
    db.commit()
    return {"ok": True}


# ---------------- users ----------------

@router.get("/users", tags=["users"])
def list_users(page: int = Query(1, ge=1), size: int = Query(20, le=200), q: str | None = None,
               db: Session = Depends(get_db), _: User = Depends(require("users", "view"))):
    stmt = select(User)
    if q:
        stmt = stmt.where(or_(User.username.ilike(f"%{q}%"), User.full_name.ilike(f"%{q}%")))
    items, total = paginate(db, stmt.order_by(User.id), page, size)
    return {"items": [user_out(u) for u in items], "total": total, "page": page, "size": size}


@router.get("/users/options", tags=["users"])
def user_options(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    """Names only, for the "entered by" (អ្នកបញ្ចូល) filter of every list."""
    return [{"id": u.id, "full_name": u.full_name, "username": u.username}
            for u in db.scalars(select(User).order_by(User.full_name))]


def _apply_user(db: Session, u: User, data: UserIn, actor: User) -> None:
    role = db.get(Role, data.role_id)
    if not role:
        raise HTTPException(400, "តួនាទីមិនត្រឹមត្រូវ")
    if role.is_super and not actor.role.is_super:
        raise HTTPException(403, "មានតែ Super Admin ទេដែលអាចកំណត់តួនាទី Super Admin")
    u.username, u.full_name, u.phone = data.username, data.full_name, data.phone
    u.role_id, u.unit_scope_id, u.active = data.role_id, data.unit_scope_id, data.active
    if data.password:
        u.password_hash = hash_password(data.password)


@router.post("/users", tags=["users"], status_code=201)
def create_user(payload: UserIn, request: Request, db: Session = Depends(get_db),
                actor: User = Depends(require("users", "create"))):
    if not payload.password:
        raise HTTPException(400, "សូមបញ្ចូលពាក្យសម្ងាត់ (យ៉ាងតិច ៨ តួ)")
    u = User()
    _apply_user(db, u, payload, actor)
    db.add(u)
    try:
        db.flush()
        audit(db, actor, "create", table="users", record_id=u.id, request=request)
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "ឈ្មោះអ្នកប្រើនេះមានរួចហើយ")
    db.refresh(u)
    return user_out(u)


@router.put("/users/{user_id}", tags=["users"])
def update_user(user_id: int, payload: UserIn, request: Request, db: Session = Depends(get_db),
                actor: User = Depends(require("users", "update"))):
    u = db.get(User, user_id)
    if not u:
        raise HTTPException(404, "រកមិនឃើញ")
    if u.role.is_super and not actor.role.is_super:
        raise HTTPException(403, "មិនអាចកែប្រែ Super Admin")
    if u.id == actor.id and not payload.active:
        raise HTTPException(400, "មិនអាចបិទគណនីខ្លួនឯងបានទេ")
    _apply_user(db, u, payload, actor)
    try:
        audit(db, actor, "update", table="users", record_id=u.id, request=request)
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "ឈ្មោះអ្នកប្រើនេះមានរួចហើយ")
    db.refresh(u)
    return user_out(u)


@router.delete("/users/{user_id}", tags=["users"])
def delete_user(user_id: int, request: Request, db: Session = Depends(get_db),
                actor: User = Depends(require("users", "delete"))):
    u = db.get(User, user_id)
    if not u:
        raise HTTPException(404, "រកមិនឃើញ")
    if u.id == actor.id or (u.role.is_super and not actor.role.is_super):
        raise HTTPException(400, "មិនអាចលុបគណនីនេះបានទេ")
    # keep history: deactivate rather than hard delete
    u.active = False
    audit(db, actor, "deactivate", table="users", record_id=u.id, request=request)
    db.commit()
    return {"ok": True}


# ---------------- roles & permissions ----------------

@router.get("/permissions/catalog", tags=["roles"])
def permission_catalog(_: User = Depends(get_current_user)):
    return {"modules": [{"key": k, "name": v, "actions": module_actions(k)} for k, v in MODULES.items()], "actions": ACTIONS}


@router.get("/roles", tags=["roles"])
def list_roles(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    roles = db.scalars(select(Role).order_by(Role.id)).all()
    return {"items": [{**to_dict(r), "permissions": sorted(p.key for p in r.permissions)} for r in roles],
            "total": len(roles)}


@router.post("/roles", tags=["roles"], status_code=201)
def create_role(payload: RoleIn, request: Request, db: Session = Depends(get_db),
                actor: User = Depends(require("roles", "create"))):
    r = Role(name=payload.name, description=payload.description)
    db.add(r)
    try:
        db.flush()
        audit(db, actor, "create", r, request=request)
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "ឈ្មោះតួនាទីមានរួចហើយ")
    return to_dict(r)


@router.put("/roles/{role_id}", tags=["roles"])
def update_role(role_id: int, payload: RoleIn, request: Request, db: Session = Depends(get_db),
                actor: User = Depends(require("roles", "update"))):
    r = db.get(Role, role_id)
    if not r:
        raise HTTPException(404, "រកមិនឃើញ")
    r.name, r.description = payload.name, payload.description
    try:
        db.flush()
        audit(db, actor, "update", r, request=request)
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "ឈ្មោះតួនាទីមានរួចហើយ")
    return to_dict(r)


@router.put("/roles/{role_id}/permissions", tags=["roles"])
def set_role_permissions(role_id: int, payload: RolePermissionsIn, request: Request, db: Session = Depends(get_db),
                         actor: User = Depends(require("roles", "update"))):
    r = db.get(Role, role_id)
    if not r:
        raise HTTPException(404, "រកមិនឃើញ")
    if r.is_super:
        raise HTTPException(400, "Super Admin មានសិទ្ធិទាំងអស់ដោយស្វ័យប្រវត្តិ")
    wanted = set(payload.permissions)
    before = {"permissions": sorted(p.key for p in r.permissions)}
    r.permissions = [p for p in db.scalars(select(Permission)) if p.key in wanted]
    audit(db, actor, "update_permissions", table="roles", record_id=r.id, before=before, request=request)
    db.commit()
    return {"ok": True, "count": len(r.permissions)}


@router.delete("/roles/{role_id}", tags=["roles"])
def delete_role(role_id: int, request: Request, db: Session = Depends(get_db),
                actor: User = Depends(require("roles", "delete"))):
    r = db.get(Role, role_id)
    if not r or r.is_super:
        raise HTTPException(400, "មិនអាចលុបតួនាទីនេះបានទេ")
    if db.scalar(select(User.id).where(User.role_id == role_id).limit(1)):
        raise HTTPException(409, "តួនាទីនេះកំពុងប្រើដោយអ្នកប្រើប្រាស់")
    db.delete(r)
    audit(db, actor, "delete", table="roles", record_id=role_id, request=request)
    db.commit()
    return {"ok": True}


# ---------------- settings ----------------

@router.get("/settings", tags=["settings"])
def get_settings(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    return {s.key: {"value": s.value, "description": s.description}
            for s in db.scalars(select(Setting)) if s.key != "logo_path"}


@router.put("/settings", tags=["settings"])
def put_settings(payload: SettingsIn, request: Request, db: Session = Depends(get_db),
                 actor: User = Depends(require("settings", "update"))):
    for k, v in payload.values.items():
        if k == "logo_path":  # set only by the logo upload: it is served without login
            continue
        s = db.get(Setting, k)
        if s:
            s.value = v
        else:
            db.add(Setting(key=k, value=v))
    audit(db, actor, "update", table="settings", record_id="*", request=request)
    db.commit()
    return {"ok": True}


# ---------------- logo (ផ្លាស់ប្ដូរ logo) ----------------

@router.get("/settings/logo", tags=["settings"])
def get_logo(db: Session = Depends(get_db)):
    """Public: the login page shows it before anyone signs in."""
    path = get_setting(db, "logo_path")
    if not path:
        raise HTTPException(404, "គ្មានឡូហ្គោ")
    return FileResponse(storage_file(path), headers={"Cache-Control": "no-cache"})


@router.post("/settings/logo", tags=["settings"])
def upload_logo(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db),
                actor: User = Depends(require("settings", "update"))):
    path, _ = save_upload(file, "branding", images_only=True)
    old = db.get(Setting, "logo_path")
    if old:
        delete_storage_file(old.value)
        old.value = path
    else:
        db.add(Setting(key="logo_path", value=path, description="Logo file"))
    audit(db, actor, "upload_logo", table="settings", record_id="logo_path", request=request)
    db.commit()
    return {"ok": True}


@router.delete("/settings/logo", tags=["settings"])
def delete_logo(request: Request, db: Session = Depends(get_db), actor: User = Depends(require("settings", "update"))):
    s = db.get(Setting, "logo_path")
    if s:
        delete_storage_file(s.value)
        db.delete(s)
        audit(db, actor, "delete", table="settings", record_id="logo_path", request=request)
        db.commit()
    return {"ok": True}


def get_setting(db: Session, key: str, default: str = "") -> str:
    s = db.get(Setting, key)
    return s.value if s and s.value is not None else default


# ---------------- audit log ----------------

@router.get("/audit-logs", tags=["audit"])
def audit_logs(page: int = Query(1, ge=1), size: int = Query(50, le=200), table: str | None = None,
               user_id: int | None = None, action: str | None = None, record_id: str | None = None,
               date_from: date | None = None, date_to: date | None = None,
               db: Session = Depends(get_db), _: User = Depends(require("audit", "view"))):
    stmt = select(AuditLog)
    if table:
        stmt = stmt.where(AuditLog.table_name == table)
    if user_id:
        stmt = stmt.where(AuditLog.user_id == user_id)
    if action:
        stmt = stmt.where(AuditLog.action == action)
    if record_id:
        stmt = stmt.where(AuditLog.record_id == record_id)
    if date_from:
        stmt = stmt.where(AuditLog.created_at >= date_from)
    if date_to:  # the whole last day
        stmt = stmt.where(AuditLog.created_at < date_to + timedelta(days=1))
    items, total = paginate(db, stmt.order_by(AuditLog.id.desc()), page, size)
    return {
        "items": [{**to_dict(a), "username": a.user.username if a.user else None} for a in items],
        "total": total, "page": page, "size": size,
    }


# ---------------- saved filters & export by filter ----------------

def saved_filter_out(f: SavedFilter, user: User) -> dict:
    return {"id": f.id, "module": f.module, "name": f.name, "params": f.params, "shared": f.shared,
            "owner": f.user.full_name if f.user else None, "mine": f.user_id == user.id}


@router.get("/saved-filters", tags=["saved filters"])
def list_saved_filters(module: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """The user's own filters of a list, then the ones other users shared."""
    rows = db.scalars(select(SavedFilter).where(SavedFilter.module == module,
                                                or_(SavedFilter.user_id == user.id, SavedFilter.shared.is_(True)))
                      .order_by(SavedFilter.name))
    return sorted((saved_filter_out(f, user) for f in rows), key=lambda x: not x["mine"])


@router.post("/saved-filters", tags=["saved filters"], status_code=201)
def save_filter(payload: SavedFilterIn, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """Saving under a name the user already used for this list replaces that filter."""
    params = {k: v for k, v in payload.params.items() if v not in (None, "", [])}
    f = db.scalar(select(SavedFilter).where(SavedFilter.user_id == user.id, SavedFilter.module == payload.module,
                                            SavedFilter.name == payload.name))
    if not f:
        f = SavedFilter(user_id=user.id, module=payload.module, name=payload.name)
        db.add(f)
    f.params, f.shared = params, payload.shared
    db.commit()
    db.refresh(f)
    return saved_filter_out(f, user)


@router.delete("/saved-filters/{filter_id}", tags=["saved filters"])
def delete_saved_filter(filter_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    f = db.get(SavedFilter, filter_id)
    if not f:
        raise HTTPException(404, "រកមិនឃើញ")
    if f.user_id != user.id and not user.role.is_super:
        raise HTTPException(403, "អាចលុបបានតែ Filter របស់ខ្លួនឯងប៉ុណ្ណោះ")
    db.delete(f)
    db.commit()
    return {"ok": True}


@router.post("/export/table", tags=["export"])
def export_table(payload: ExportTableIn, request: Request, db: Session = Depends(get_db),
                 user: User = Depends(get_current_user)):
    """Excel of a list exactly as filtered on screen. The page sends the rows it loaded through the module's own
    (permission- and scope-checked) list endpoint; this only needs the module's export permission."""
    if not user.can(payload.module, "export"):
        raise HTTPException(403, "គ្មានសិទ្ធិនាំចេញ")
    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    ncol = len(payload.columns) + 1
    ws.append([payload.title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ncol, 1))
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.append(["ល.រ", *payload.columns])
    for c in ws[2]:
        c.font = Font(bold=True)
    for i, row in enumerate(payload.rows, 1):
        ws.append([i, *row])
    for col in ws.iter_cols(min_row=2, max_row=2):
        ws.column_dimensions[col[0].column_letter].width = 8 if col[0].column == 1 else 22
    ws.freeze_panes = "A3"
    audit(db, user, "export", table=payload.module, record_id="*", request=request,
          before={"title": payload.title, "rows": len(payload.rows)})
    db.commit()
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=export.xlsx"})

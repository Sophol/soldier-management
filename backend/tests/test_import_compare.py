"""Excel import of professional soldiers and comparison with an Excel list (New.txt §2, §3)."""

from io import BytesIO

from openpyxl import Workbook, load_workbook

from tests.test_flow import API, _lookups, _soldier, _units


def _xlsx(header, *rows):
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(list(r))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(c, path, content, **form):
    data = {k: str(v).lower() if isinstance(v, bool) else v for k, v in form.items()}
    return c.post(f"{API}/soldiers/{path}", files={"file": ("list.xlsx", content)}, data=data)


def test_template_has_every_column_and_value_lists(admin):
    r = admin.get(f"{API}/soldiers/import/template")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    header = [c.value for c in wb.worksheets[0][1]]
    assert header[0] == "អត្តលេខ *" and "បញ្ចូលក្របខ័ណ្ឌលេខ" in header
    assert wb.worksheets[1].title == "តម្លៃដែលអាចប្រើ"


def test_import_dry_run_then_apply(admin):
    lk = _lookups(admin)
    root, unit, _ = _units(admin)
    rank = lk["ranks"][-3]["name_kh"]
    path = f"{root['name_kh']} › {unit['name_kh']}"
    header = ["អត្តលេខ *", "គោត្តនាមនិងនាម *", "ភេទ *", "ថ្ងៃខែឆ្នាំកំណើត", "ឋានន្តរស័ក្តិ", "ស្ថាប័នកងឯកភាព-អង្គភាព",
              "ថ្ងៃខែឆ្នាំបម្រើចូលទ័ព", "បញ្ចូលក្របខ័ណ្ឌលេខ", "ជួរដែលមិនស្គាល់"]
    content = _xlsx(header,
                    ["TEST-X1", "សុខ ដារ៉ា", "ប្រុស", "០១/០៥/១៩៩០", rank, path, "15/01/2010", "ក.123", "x"],
                    ["TEST-X2", "ចាន់ ធីតា", "ស្រី", "1992-03-04", rank, unit["name_kh"], None, None, None],
                    ["TEST-X1", "duplicate", "ប្រុស", None, None, None, None, None, None],
                    ["TEST-X3", "x", "?", "32/13/2000", "no such rank", "no unit", None, None, None])
    plan = _post(admin, "import", content).json()
    assert plan["dry_run"] and (plan["create"], plan["error"]) == (2, 2)
    assert plan["unknown_columns"] == ["ជួរដែលមិនស្គាល់"]
    bad = next(r for r in plan["rows"] if r["military_id"] == "TEST-X3")
    assert len(bad["errors"]) == 4  # gender, date, rank, unit
    assert admin.get(f"{API}/soldiers", params={"q": "TEST-X1"}).json()["total"] == 0  # nothing written
    done = _post(admin, "import", content, dry_run=False).json()
    assert (done["create"], done["error"]) == (2, 2)
    s = admin.get(f"{API}/soldiers/{next(r['soldier_id'] for r in done['rows'] if r['military_id'] == 'TEST-X1')}").json()
    assert (s["date_of_birth"], s["establishment_no"], s["current_unit_id"]) == ("1990-05-01", "ក.123", unit["id"])
    mv = admin.get(f"{API}/strength-movements", params={"soldier_id": s["id"]}).json()
    assert mv["total"] == 1 and mv["items"][0]["reason"] == "recruit"
    # importing again: existing IDs are refused unless updating is asked for; updates keep blank cells
    again = _post(admin, "import", _xlsx(["អត្តលេខ", "លេខទូរសព្ទ"], ["TEST-X1", "012 345 678"]), dry_run=False).json()
    assert again["error"] == 1
    upd = _post(admin, "import", _xlsx(["អត្តលេខ", "លេខទូរសព្ទ"], ["TEST-X1", "012 345 678"]),
                dry_run=False, update_existing=True).json()
    assert upd["update"] == 1
    s2 = admin.get(f"{API}/soldiers/{s['id']}").json()
    assert s2["phone"] == "012 345 678" and s2["establishment_no"] == "ក.123"


def test_export_round_trips_into_compare(admin):
    lk = _lookups(admin)
    _, unit, other = _units(admin)
    a = _soldier(admin, lk, unit["id"], mid="TEST-C1")
    b = _soldier(admin, lk, unit["id"], mid="TEST-C2", national_id=None)
    exported = admin.get(f"{API}/soldiers/export", params={"unit_id": unit["id"]}).content
    same = _post(admin, "compare", exported, unit_id=unit["id"]).json()
    assert same["same"] == 2 and same["different"] == same["missing_in_file"] == 0
    # change the file: a phone number differs, one soldier is missing, one is unknown
    header = ["អត្តលេខ", "គោត្តនាមនិងនាម", "លេខទូរសព្ទ", "ស្ថាប័នកងឯកភាព-អង្គភាព"]
    content = _xlsx(header, ["TEST-C1", a["name_kh"], "099 999 999", unit["name_kh"]], ["TEST-ZZ", "x", None, None])
    res = _post(admin, "compare", content, unit_id=unit["id"], fields="phone,current_unit_id").json()
    by_mid = {r["military_id"]: r for r in res["rows"]}
    assert by_mid["TEST-C1"]["status"] == "different" and [d["field"] for d in by_mid["TEST-C1"]["diffs"]] == ["phone"]
    assert by_mid["TEST-ZZ"]["status"] == "missing_in_db" and by_mid[b["military_id"]]["status"] == "missing_in_file"
    xlsx = _post(admin, "compare", content, unit_id=unit["id"], format="xlsx")
    assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK"
    assert other


def test_khmer_digits_in_names_and_dates(admin):
    """Unit names keep their Khmer digits (លេខ១), dates accept them (០១/០៥/១៩៩០)."""
    lk = _lookups(admin)
    root = admin.post(f"{API}/units", json={"name_kh": "ក្រសួង TEST", "level": 1}).json()
    unit = admin.post(f"{API}/units", json={"name_kh": "វរសេនាតូចលេខ១ TEST", "parent_id": root["id"], "level": 1}).json()
    content = _xlsx(["អត្តលេខ", "គោត្តនាមនិងនាម", "ភេទ", "ថ្ងៃខែឆ្នាំកំណើត", "ស្ថាប័នកងឯកភាព-អង្គភាព"],
                    ["TEST-KD1", "x", "ប្រុស", "០១/០៥/១៩៩០", f"{root['name_kh']} › {unit['name_kh']}"])
    plan = _post(admin, "import", content).json()
    assert plan["create"] == 1, plan["rows"]
    assert lk
